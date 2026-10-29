"""Excel workbook and plain-text summary."""
from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .kpis import Kpis

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)


def _fmt_value(label: str, value) -> tuple[object, str | None]:
    if value is None:
        return "n/a", None
    if isinstance(value, float):
        lowered = label.lower()
        if "rate" in lowered or "share" in lowered:
            return value, "0.0%"
        return round(value, 1), "0.0"
    return value, None


def _write_table(ws, df: pd.DataFrame, top: int, left: int = 1) -> int:
    for j, col in enumerate(df.columns):
        cell = ws.cell(row=top, column=left + j, value=str(col).replace("_", " ").title())
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, value in enumerate(row):
            value = None if pd.isna(value) else (value.to_pydatetime() if isinstance(value, pd.Timestamp) else value)
            cell = ws.cell(row=top + i, column=left + j, value=value)
            name = str(df.columns[j])
            if name in ("share", "completion_rate"):
                cell.number_format = "0.0%"
            elif name == "average_score":
                cell.number_format = "0.0"
            elif hasattr(value, "year"):
                cell.number_format = "yyyy-mm-dd"
    return top + len(df)


def _autosize(ws) -> None:
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 10), 60)


def _add_bar(ws, title: str, rows: int, anchor: str, cat_col: int = 1, val_col: int = 2, header_row: int = 1):
    if rows < 1:
        return
    chart = BarChart()
    chart.title, chart.legend = title, None
    chart.add_data(Reference(ws, min_col=val_col, min_row=header_row, max_row=header_row + rows), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=cat_col, min_row=header_row + 1, max_row=header_row + rows))
    chart.height, chart.width = 7.5, 14
    ws.add_chart(chart, anchor)


def build_workbook(kpis: Kpis, quality: list[dict], period_label: str, entity: str, path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"{entity} - AML/CFT management report"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Period: {period_label} ({kpis.period_start:%Y-%m-%d} to {kpis.period_end:%Y-%m-%d})"
    ws["A4"], ws["B4"] = "Indicator", "Value"
    for c in (ws["A4"], ws["B4"]):
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    for i, (label, value) in enumerate(kpis.headline.items(), start=5):
        shown, fmt = _fmt_value(label, value)
        ws.cell(row=i, column=1, value=label)
        cell = ws.cell(row=i, column=2, value=shown)
        cell.alignment = Alignment(horizontal="right")
        if fmt:
            cell.number_format = fmt
    note_row = 6 + len(kpis.headline)
    ws.cell(row=note_row, column=1, value=f"Data-quality findings: {len(quality)} (see 'Data Quality' sheet)")
    _autosize(ws)

    def sheet(name: str, df: pd.DataFrame | None, chart: str | None = None, pie: bool = False):
        if df is None or df.empty:
            return
        s = wb.create_sheet(name)
        last = _write_table(s, df, top=1)
        if chart:
            if pie:
                pc = PieChart()
                pc.title = chart
                pc.add_data(Reference(s, min_col=2, min_row=1, max_row=last), titles_from_data=True)
                pc.set_categories(Reference(s, min_col=1, min_row=2, max_row=last))
                pc.height, pc.width = 7.5, 10
                s.add_chart(pc, "F2")
            else:
                _add_bar(s, chart, rows=last - 1, anchor="F2")
        _autosize(s)

    sheet("Alerts by Rule", kpis.alerts_by_rule, "Alerts by rule")
    sheet("Alerts by Severity", kpis.alerts_by_severity, "Alerts by severity")
    sheet("Alert Outcomes", kpis.alerts_by_outcome)
    sheet("Customer Risk", kpis.customer_risk, "Customers by risk level", pie=True)
    sheet("Screening", kpis.screening, "Screening results")
    sheet("Training", kpis.training)

    dq = wb.create_sheet("Data Quality")
    if quality:
        _write_table(dq, pd.DataFrame(quality), top=1)
    else:
        dq["A1"] = "No data-quality issues found."
    _autosize(dq)
    wb.save(path)


def build_summary(kpis: Kpis, quality: list[dict], period_label: str, entity: str) -> str:
    lines = [f"# {entity}: AML/CFT report, {period_label}", "",
             f"Reporting period: {kpis.period_start:%Y-%m-%d} to {kpis.period_end:%Y-%m-%d}", "",
             "## Key indicators", ""]
    for label, value in kpis.headline.items():
        shown, fmt = _fmt_value(label, value)
        if fmt == "0.0%":
            shown = f"{shown:.1%}"
        lines.append(f"- {label}: **{shown}**")
    lines += ["", "## Data quality", ""]
    if quality:
        lines += [f"- {q['dataset']}: {q['check']} ({q['records']} record(s)) - {q['detail']}" for q in quality]
    else:
        lines.append("- No issues found.")
    return "\n".join(lines) + "\n"
