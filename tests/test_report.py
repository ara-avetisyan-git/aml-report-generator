from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from aml_reports import quality
from aml_reports.cli import main
from aml_reports.kpis import build_kpis
from aml_reports.loaders import read_table
from aml_reports.period import parse_period

DATA = Path(__file__).parent / "data"


def test_period_parsing():
    assert parse_period("2026-Q3") == (pd.Timestamp("2026-07-01"), pd.Timestamp("2026-09-30"))
    assert parse_period("2026-02")[1] == pd.Timestamp("2026-02-28")
    assert parse_period("2026")[0] == pd.Timestamp("2026-01-01")
    with pytest.raises(ValueError):
        parse_period("Q3 2026")


def _alerts(rows):
    cols = ["alert_id", "rule_id", "severity", "customer_id", "opened_date", "closed_date", "status", "outcome"]
    df = pd.DataFrame(rows, columns=cols)
    for c in ("opened_date", "closed_date"):
        df[c] = pd.to_datetime(df[c].replace("", pd.NA))
    return df


def test_alert_kpis_exact_values():
    alerts = _alerts([
        ["1", "r1", "high", "C1", "2026-07-01", "2026-07-05", "closed", "false_positive"],
        ["2", "r1", "high", "C2", "2026-07-02", "2026-07-12", "closed", "str_filed"],
        ["3", "r2", "low", "C3", "2026-07-03", "", "open", ""],
        ["4", "r2", "low", "C4", "2026-06-15", "2026-06-20", "closed", "no_action"],   # outside period
    ])
    k = build_kpis(pd.Timestamp("2026-07-01"), pd.Timestamp("2026-09-30"), alerts=alerts)
    h = k.headline
    assert h["Alerts opened in period"] == 3
    assert h["Alerts closed"] == 2
    assert h["Average days to close"] == 7.0          # (4 + 10) / 2
    assert h["False-positive rate (closed)"] == 0.5
    assert h["STRs filed"] == 1
    assert h["Open alerts older than 30 days"] == 1
    assert int(k.alerts_by_rule.set_index("rule_id").loc["r1", "count"]) == 2


def test_quality_checks_find_planted_problems():
    alerts = read_table("alerts", DATA / "alerts.csv")
    customers = read_table("customers", DATA / "customers.csv")
    checks = {i["check"] for i in quality.check_alerts(alerts) + quality.check_cross(alerts, customers)}
    assert "duplicate alert_id" in checks
    assert "closed alert without closed_date" in checks
    assert "closed_date before opened_date" in checks
    assert "alert for unknown customer" in checks


def test_missing_columns_are_reported(tmp_path):
    bad = tmp_path / "a.csv"
    bad.write_text("alert_id,rule_id\n1,x\n", encoding="utf-8")
    with pytest.raises(ValueError, match="missing required columns"):
        read_table("alerts", bad)


def test_end_to_end_workbook(tmp_path, capsys):
    code = main(["build", "--period", "2026-Q3", "--entity", "Demo Co",
                 "--alerts", str(DATA / "alerts.csv"), "--customers", str(DATA / "customers.csv"),
                 "--screening", str(DATA / "screening.csv"), "--training", str(DATA / "training.csv"),
                 "--output-dir", str(tmp_path)])
    assert code == 0
    wb = load_workbook(tmp_path / "aml_report_2026-Q3.xlsx")
    assert {"Summary", "Alerts by Rule", "Customer Risk", "Screening", "Training", "Data Quality"} <= set(wb.sheetnames)
    summary = wb["Summary"]
    labels = [summary.cell(row=r, column=1).value for r in range(1, 30)]
    assert "Alerts opened in period" in labels and "Training completion rate" in labels
    text = (tmp_path / "aml_report_2026-Q3.md").read_text(encoding="utf-8")
    assert "Demo Co" in text and "Data quality" in text


def test_cli_requires_input(tmp_path, capsys):
    assert main(["build", "--period", "2026-Q3", "--output-dir", str(tmp_path)]) == 2
